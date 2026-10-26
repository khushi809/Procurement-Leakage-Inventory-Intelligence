
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter

PROJECT_DIR = Path(__file__).resolve().parent.parent
OUTPUT_DIR = PROJECT_DIR / "outputs"

seller_file = OUTPUT_DIR / "seller_performance.csv"
orders_file = OUTPUT_DIR / "order_analysis.csv"
risk_file = OUTPUT_DIR / "seller_risk_analysis.csv"
excel_file = OUTPUT_DIR / "business_performance_report.xlsx"

# Load existing analysis results
sellers = pd.read_csv(seller_file)
orders = pd.read_csv(orders_file)
risk = pd.read_csv(risk_file)

# Business summary
summary = pd.DataFrame({
    "Metric": [
        "Total orders",
        "Delivered orders",
        "Item sales value (R$)",
        "Freight charges (R$)",
        "Recorded payment value (R$)",
        "Average delivery time (days)",
        "Late deliveries",
        "Sellers assessed for risk",
        "High-priority investigations",
        "Sellers requiring review",
    ],
    "Value": [
        len(orders),
        (orders["order_status"] == "delivered").sum(),
        orders["item_sales"].sum(),
        orders["freight_cost"].sum(),
        orders["payment_value"].sum(),
        orders.loc[
            orders["order_status"] == "delivered",
            "delivery_days"
        ].mean(),
        (
            (orders["order_status"] == "delivered")
            & (orders["delivery_delay_days"] > 0)
        ).sum(),
        len(risk),
        (risk["investigation_category"]
         == "High priority investigation").sum(),
        (risk["investigation_category"]
         == "Review recommended").sum(),
    ],
})

# Create workbook with three sheets
with pd.ExcelWriter(excel_file, engine="openpyxl") as writer:
    summary.to_excel(
        writer, sheet_name="Business Summary", index=False
    )
    sellers.to_excel(
        writer, sheet_name="Seller Performance", index=False
    )
    risk.to_excel(
        writer, sheet_name="Seller Risk Analysis", index=False
    )

# Apply professional formatting
workbook = load_workbook(excel_file)

for worksheet in workbook.worksheets:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False

    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            "solid", fgColor="243447"
        )
        cell.alignment = Alignment(
            horizontal="center", vertical="center",
            wrap_text=True
        )

    worksheet.row_dimensions[1].height = 30

    for column_cells in worksheet.columns:
        max_length = max(
            len(str(cell.value or ""))
            for cell in column_cells
        )
        column_letter = get_column_letter(
            column_cells[0].column
        )
        worksheet.column_dimensions[
            column_letter
        ].width = min(max_length + 3, 35)

# Conditional formatting: seller freight percentage
seller_sheet = workbook["Seller Performance"]
seller_headers = {
    cell.value: cell.column for cell in seller_sheet[1]
}
freight_col = seller_headers.get("freight_pct_of_sales")

if freight_col and seller_sheet.max_row >= 2:
    col_letter = get_column_letter(freight_col)
    seller_sheet.conditional_formatting.add(
        f"{col_letter}2:{col_letter}{seller_sheet.max_row}",
        ColorScaleRule(
            start_type="min", start_color="63BE7B",
            mid_type="percentile", mid_value=50,
            mid_color="FFEB84",
            end_type="max", end_color="F8696B",
        ),
    )

# Conditional formatting: seller risk flags
risk_sheet = workbook["Seller Risk Analysis"]
risk_headers = {
    cell.value: cell.column for cell in risk_sheet[1]
}

for header in ["freight_to_sales_pct", "late_delivery_pct"]:
    col = risk_headers.get(header)
    if col and risk_sheet.max_row >= 2:
        letter = get_column_letter(col)
        risk_sheet.conditional_formatting.add(
            f"{letter}2:{letter}{risk_sheet.max_row}",
            ColorScaleRule(
                start_type="min", start_color="63BE7B",
                mid_type="percentile", mid_value=50,
                mid_color="FFEB84",
                end_type="max", end_color="F8696B",
            ),
        )

category_col = risk_headers.get("investigation_category")

if category_col and risk_sheet.max_row >= 2:
    for row in range(2, risk_sheet.max_row + 1):
        cell = risk_sheet.cell(row=row, column=category_col)
        if cell.value == "High priority investigation":
            cell.fill = PatternFill(
                "solid", fgColor="F4CCCC"
            )
        elif cell.value == "Review recommended":
            cell.fill = PatternFill(
                "solid", fgColor="FFF2CC"
            )
        elif cell.value == "Monitor - limited data":
            cell.fill = PatternFill(
                "solid", fgColor="D9EAF7"
            )

workbook.save(excel_file)

print("Excel report updated successfully!")
print(f"Sheets: {workbook.sheetnames}")
print(f"Saved to: {excel_file}")
